#!/usr/bin/env python3
"""Fold the tiering output into XP_AE_Rebalance_Model.xlsx.

Account-level tiers are written as static values on Assignments, because a tier is a
property of the account rather than of who is assigned to it. The XP Check rollups are
written as formulas, so the tier mix recomputes when someone changes the Revised XP
dropdown. Adds a Growth Tiers sheet carrying all 51 territories.

Idempotent: re-run after build_tiering.py.

Run:  python3 tiering/build_tiering.py && python3 tiering/update_workbook.py
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
BOOK = ROOT / "XP_AE_Rebalance_Model.xlsx"
LAST = 278  # last populated Assignments row, matching the existing formulas

HEAD_FILL = PatternFill("solid", fgColor="1B2A33")
HEAD_FONT = Font(color="F6F8F9", bold=True)
TIER_FILL = {
    1: PatternFill("solid", fgColor="D8EDE2"),
    2: PatternFill("solid", fgColor="DAEAF3"),
    3: PatternFill("solid", fgColor="F6E2D6"),
    4: PatternFill("solid", fgColor="E7EAEC"),
}

ASSIGN_COLS = [
    ("Territory Tier", 13), ("Territory Score", 15), ("Capability Whitespace", 20),
    ("Growth Potential", 17), ("Growth Tier", 12), ("Growth Tier Label", 18),
    ("Growth Load", 12), ("Capability Data", 15),
]
CHECK_COLS = [
    ("Tier 1 Growth", 13), ("Tier 2 Expansion", 15), ("Tier 3 Defend", 13),
    ("Tier 4 Maintain", 15), ("Growth Share", 13), ("Growth Load", 12),
    ("Defend ARR", 13),
]


def header(ws, row, col, text, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.fill, c.font = HEAD_FILL, HEAD_FONT
    c.alignment = Alignment(wrap_text=True, vertical="center")
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def main():
    t = json.loads((ROOT / "tiering" / "tiering.json").read_text())
    by_key = {(a["account"], a["state"]): a for a in t["accounts"]}
    load = {int(k): v for k, v in t["weights"]["growth_load"].items()}

    wb = openpyxl.load_workbook(BOOK)

    # ---------------------------------------------------------------- Assignments
    ws = wb["Assignments"]
    first = 24  # column X, leaving the complex-threshold parameters in V and W alone
    for i, (name, width) in enumerate(ASSIGN_COLS):
        header(ws, 1, first + i, name, width)

    for row in range(2, LAST + 1):
        name = ws.cell(row=row, column=1).value
        if not name:
            continue
        a = by_key.get((name, (ws.cell(row=row, column=3).value or "").strip()))
        if not a or a["tier"] is None:
            continue
        vals = [a["territory_tier"], a["territory_score"], a["account_whitespace"],
                a["growth_potential"], a["tier"], a["tier_label"],
                load[a["tier"]] if a["countable"] else 0, a["capability_data"]]
        for i, v in enumerate(vals):
            c = ws.cell(row=row, column=first + i, value=v)
            if i == 4:
                c.fill = TIER_FILL[a["tier"]]
        ws.cell(row=row, column=first + 2).number_format = "0.00"
        ws.cell(row=row, column=first + 3).number_format = "0.000"
    ws.auto_filter.ref = f"A1:{get_column_letter(first + len(ASSIGN_COLS) - 1)}{LAST}"

    # ---------------------------------------------------------------- XP Check
    ck = wb["XP Check"]
    cfirst = 20  # column T, clear of the cap parameters in Q and R
    for i, (name, width) in enumerate(CHECK_COLS):
        header(ck, 1, cfirst + i, name, width)

    A = "Assignments!"
    xp = f"{A}$G$2:$G${LAST}"
    cap = f"{A}$N$2:$N${LAST}"
    tier = f"{A}$AB$2:$AB${LAST}"
    for row in range(2, 15):
        if not ck.cell(row=row, column=1).value:
            continue
        me = f"$A{row}"
        for k in range(1, 5):
            ck.cell(row=row, column=cfirst + k - 1,
                    value=f'=COUNTIFS({xp},{me},{cap},"Yes",{tier},{k})')
        t1, t2 = get_column_letter(cfirst), get_column_letter(cfirst + 1)
        tot = f'COUNTIFS({xp},{me},{cap},"Yes")'
        ck.cell(row=row, column=cfirst + 4,
                value=f'=IFERROR(({t1}{row}+{t2}{row})/{tot},"")').number_format = "0%"
        ck.cell(row=row, column=cfirst + 5,
                value=f'=SUMPRODUCT(({xp}={me})*({cap}="Yes")*{A}$AD$2:$AD${LAST})'
                ).number_format = "0.0"
        ck.cell(row=row, column=cfirst + 6,
                value=f'=SUMIFS({A}$M$2:$M${LAST},{xp},{me},{cap},"Yes",{tier},3)'
                ).number_format = '$#,##0'

    # ---------------------------------------------------------------- Growth Tiers
    if "Growth Tiers" in wb.sheetnames:
        del wb["Growth Tiers"]
    gt = wb.create_sheet("Growth Tiers", wb.sheetnames.index("Lookups"))
    cols = [
        ("State", 20), ("State-agency tier", 11), ("State-agency score", 11),
        ("State accts in book", 10), ("Local ENT tier", 10), ("Local ENT score", 10),
        ("Local accts in book", 10), ("Tiers disagree", 10),
        ("Jurisdictions 100k+", 11), ("Held in this book", 11),
        ("Fiscal posture", 13), ("Rainy day % of spend", 11), ("Budget cycle", 11),
        ("IT modernization signal", 15), ("Grant environment", 13),
        ("ARPA cliff exposure", 13), ("Purchasing access", 14),
        ("League channel taken", 11), ("Local capacity", 13), ("Property tax cap", 11),
        ("Population", 12), ("Pop % chg", 10), ("Domestic mig / 1k", 11),
        ("Research confidence", 11), ("State ARR", 13), ("Local ARR", 13),
    ]
    for i, (name, width) in enumerate(cols, start=1):
        header(gt, 1, i, name, width)
    gt.freeze_panes = "B2"

    ordered = sorted(t["territories"].values(),
                     key=lambda r: (-(r["in_book_state_accounts"] + r["in_book_local_accounts"]) > 0,
                                    -max(r["state_side_score"], r["local_side_score"])))
    for row, r in enumerate(ordered, start=2):
        vals = [
            r["state"], r["state_side_tier"], r["state_side_score"],
            r["in_book_state_accounts"], r["local_side_tier"], r["local_side_score"],
            r["in_book_local_accounts"], "yes" if r["split_territory"] else "",
            r["addressable_over_100k"], r["local_accounts_held"],
            r["posture"], r["rdf_pct"], r["budget_cycle"], r["it_mod_signal"],
            r["grant_env"], r["arpa_cliff"], r["purchasing_access"],
            "yes" if r["channel_occupied"] else "", r["local_capacity"],
            "yes" if r["tax_cap_pressure"] else "", r["population"], r["pop_pct_chg"],
            r["dom_mig_per_1k"], r["confidence"], r["state_arr"], r["local_arr"],
        ]
        for i, v in enumerate(vals, start=1):
            gt.cell(row=row, column=i, value=v)
        gt.cell(row=row, column=21).number_format = "#,##0"
        gt.cell(row=row, column=25).number_format = '$#,##0'
        gt.cell(row=row, column=26).number_format = '$#,##0'
    gt.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(ordered) + 1}"

    # ---------------------------------------------------------------- Read Me
    rm = wb["Read Me"]
    marker = "Growth tiering (added)"
    existing = next((r for r in range(1, rm.max_row + 1)
                     if rm.cell(row=r, column=1).value == marker), None)
    if existing:
        # Re-runs replace the block rather than stacking another copy underneath it.
        rm.delete_rows(existing - 1, rm.max_row - existing + 2)
    start = rm.max_row + 2
    notes = [
        marker,
        "",
        "Each account carries a growth tier on Assignments columns X to AE. The tier crosses growth",
        "potential - 60% the territory score, 40% the capability whitespace left in the account -",
        "against ARR, splitting both at the median of the 209 countable accounts.",
        "",
        "  Tier 1 Growth engine   high potential, high ARR   expansion motion",
        "  Tier 2 Expansion       high potential, lower ARR  land and expand",
        "  Tier 3 Defend          lower potential, high ARR  retention and adoption",
        "  Tier 4 Maintain        lower potential, lower ARR efficient or pooled coverage",
        "",
        "XP Check columns T to Z roll the tiers up per XP and recompute when the Revised XP",
        "dropdown changes. Growth load weights the countable book 1.0 / 0.8 / 0.5 / 0.3 by tier.",
        "",
        "The Growth Tiers sheet carries all 50 states plus DC, scored separately for the",
        "state-agency motion and the Local ENT motion. They disagree in 32 of 51 states.",
        "",
        "Caveat: 'Held in this book' counts only Enterprise, so whitespace is an upper bound.",
        "Rebuild with: python3 tiering/build_tiering.py && python3 tiering/update_workbook.py",
        "Method, sources and full caveats: research/ENT-book-growth-tiering-hypothesis.md",
    ]
    for i, line in enumerate(notes):
        c = rm.cell(row=start + i, column=1, value=line)
        if i == 0:
            c.font = Font(bold=True)

    wb.save(BOOK)
    print(f"workbook updated: {len(ordered)} territories, "
          f"{sum(1 for a in t['accounts'] if a['tier'])} accounts tiered")


if __name__ == "__main__":
    main()
